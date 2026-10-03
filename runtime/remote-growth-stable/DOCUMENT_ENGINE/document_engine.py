from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import json
import os
import re
import shutil
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
import pymupdf
import xlrd

PROD_ROOT = Path(__file__).resolve().parent
BACKUP_ROOT = PROD_ROOT / "backups"
SUPPORTED = {".docx", ".pdf", ".xlsx", ".xls", ".pptx", ".csv", ".txt", ".md", ".json", ".jsonl", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf"}
EDITABLE = {".docx", ".xlsx", ".pptx", ".csv", ".txt", ".md", ".json", ".jsonl", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf"}
TEXTLIKE = {".csv", ".txt", ".md", ".json", ".jsonl", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf"}

class EngineError(Exception):
    pass

def emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=True, default=str))

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def stat_base(path: Path) -> dict[str, Any]:
    st = path.stat()
    return {
        "path": str(path),
        "name": path.name,
        "extension": path.suffix.lower(),
        "size": st.st_size,
        "modified": datetime.fromtimestamp(st.st_mtime).astimezone().isoformat(timespec="seconds"),
        "sha256": sha256(path),
    }

def ensure_path(raw: str) -> Path:
    p = Path(raw).expanduser()
    if not p.exists():
        raise EngineError(f"File not found: {p}")
    if not p.is_file():
        raise EngineError(f"Not a file: {p}")
    if p.suffix.lower() not in SUPPORTED:
        raise EngineError(f"Unsupported extension: {p.suffix.lower()}")
    return p

def read_textlike(path: Path) -> tuple[str, str]:
    data = path.read_bytes()
    if b"\x00" in data[:8192]:
        raise EngineError("File appears binary")
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replace"

def docx_blocks(path: Path) -> list[dict[str, Any]]:
    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    for i, p in enumerate(doc.paragraphs, 1):
        if p.text.strip():
            blocks.append({"kind": "paragraph", "index": i, "text": p.text})
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            vals = [cell.text for cell in row.cells]
            blocks.append({"kind": "table_row", "table": ti, "row": ri, "text": " | ".join(vals)})
    for si, sec in enumerate(doc.sections, 1):
        for kind, paragraphs in (
            ("header", sec.header.paragraphs),
            ("footer", sec.footer.paragraphs),
        ):
            for pi, p in enumerate(paragraphs, 1):
                if p.text.strip():
                    blocks.append({"kind": kind, "section": si, "index": pi, "text": p.text})
    return blocks

def pdf_blocks(path: Path) -> list[dict[str, Any]]:
    doc = pymupdf.open(str(path))
    blocks = []
    try:
        for i, page in enumerate(doc, 1):
            blocks.append({"kind": "page", "page": i, "text": page.get_text("text") or ""})
    finally:
        doc.close()
    return blocks

def xlsx_blocks(path: Path, include_formulas: bool = True) -> list[dict[str, Any]]:
    wb = load_workbook(str(path), read_only=True, data_only=not include_formulas)
    blocks = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                vals = []
                coords = []
                for cell in row:
                    if cell.value is not None:
                        vals.append(str(cell.value))
                        coords.append(cell.coordinate)
                if vals:
                    blocks.append({
                        "kind": "sheet_row",
                        "sheet": ws.title,
                        "row": row[0].row if row else None,
                        "cells": coords,
                        "text": " | ".join(vals),
                    })
    finally:
        wb.close()
    return blocks

def xls_blocks(path: Path) -> list[dict[str, Any]]:
    book = xlrd.open_workbook(str(path), on_demand=True)
    blocks = []
    try:
        for sheet in book.sheets():
            for r in range(sheet.nrows):
                vals = []
                for c in range(sheet.ncols):
                    v = sheet.cell_value(r, c)
                    if v not in ("", None):
                        vals.append(str(v))
                if vals:
                    blocks.append({
                        "kind": "sheet_row",
                        "sheet": sheet.name,
                        "row": r + 1,
                        "text": " | ".join(vals),
                    })
    finally:
        book.release_resources()
    return blocks

def pptx_blocks(path: Path) -> list[dict[str, Any]]:
    prs = Presentation(str(path))
    blocks = []
    for si, slide in enumerate(prs.slides, 1):
        for shi, shape in enumerate(slide.shapes, 1):
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                blocks.append({"kind": "slide_text", "slide": si, "shape": shi, "text": shape.text})
            if getattr(shape, "has_table", False):
                for ri, row in enumerate(shape.table.rows, 1):
                    vals = [cell.text for cell in row.cells]
                    blocks.append({"kind": "slide_table_row", "slide": si, "shape": shi, "row": ri, "text": " | ".join(vals)})
    return blocks

def csv_blocks(path: Path) -> list[dict[str, Any]]:
    text, enc = read_textlike(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(text.splitlines(), dialect)
    blocks = []
    for i, row in enumerate(reader, 1):
        blocks.append({"kind": "csv_row", "row": i, "text": " | ".join(row)})
    return blocks

def text_blocks(path: Path) -> list[dict[str, Any]]:
    text, _ = read_textlike(path)
    return [{"kind": "line", "line": i, "text": line} for i, line in enumerate(text.splitlines(), 1)]

def blocks_for(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".docx":
        return docx_blocks(path)
    if ext == ".pdf":
        return pdf_blocks(path)
    if ext == ".xlsx":
        return xlsx_blocks(path)
    if ext == ".xls":
        return xls_blocks(path)
    if ext == ".pptx":
        return pptx_blocks(path)
    if ext == ".csv":
        return csv_blocks(path)
    if ext in TEXTLIKE:
        return text_blocks(path)
    raise EngineError(f"Unsupported extension: {ext}")

def plain_text(path: Path, max_chars: int | None = None) -> str:
    text = "\n".join(block.get("text", "") for block in blocks_for(path))
    return text if max_chars is None else text[:max_chars]

def inspect_file(path: Path) -> dict[str, Any]:
    ext = path.suffix.lower()
    base = stat_base(path)
    detail: dict[str, Any] = {}

    if ext == ".docx":
        doc = Document(str(path))
        blocks = docx_blocks(path)
        detail = {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "sections": len(doc.sections),
            "text_blocks": len(blocks),
            "text_chars": sum(len(x["text"]) for x in blocks),
            "words": sum(len(re.findall(r"\S+", x["text"])) for x in blocks),
        }
    elif ext == ".pdf":
        reader = PdfReader(str(path))
        blocks = pdf_blocks(path)
        detail = {
            "pages": len(reader.pages),
            "text_chars": sum(len(x["text"]) for x in blocks),
            "metadata": {str(k): str(v) for k, v in (reader.metadata or {}).items()},
        }
    elif ext == ".xlsx":
        wb = load_workbook(str(path), read_only=True, data_only=False)
        sheets = []
        formulas = 0
        nonempty = 0
        try:
            for ws in wb.worksheets:
                sheet_nonempty = 0
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            sheet_nonempty += 1
                            nonempty += 1
                            if isinstance(cell.value, str) and cell.value.startswith("="):
                                formulas += 1
                sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "nonempty_cells": sheet_nonempty})
        finally:
            wb.close()
        detail = {"sheets": sheets, "sheet_count": len(sheets), "nonempty_cells": nonempty, "formula_cells": formulas}
    elif ext == ".xls":
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            detail = {
                "sheet_count": book.nsheets,
                "sheets": [{"name": s.name, "rows": s.nrows, "columns": s.ncols} for s in book.sheets()],
                "editable": False,
            }
        finally:
            book.release_resources()
    elif ext == ".pptx":
        prs = Presentation(str(path))
        blocks = pptx_blocks(path)
        detail = {"slides": len(prs.slides), "text_blocks": len(blocks), "text_chars": sum(len(x["text"]) for x in blocks)}
    elif ext == ".csv":
        blocks = csv_blocks(path)
        detail = {"rows": len(blocks), "text_chars": sum(len(x["text"]) for x in blocks)}
    else:
        text, enc = read_textlike(path)
        detail = {"encoding": enc, "lines": len(text.splitlines()), "text_chars": len(text)}

    return {"ok": True, "command": "inspect", **base, "editable": ext in EDITABLE, "detail": detail}

def filter_blocks(blocks: list[dict[str, Any]], page: int | None, sheet: str | None, slide: int | None) -> list[dict[str, Any]]:
    out = blocks
    if page is not None:
        out = [x for x in out if x.get("page") == page]
    if sheet is not None:
        out = [x for x in out if str(x.get("sheet", "")).casefold() == sheet.casefold()]
    if slide is not None:
        out = [x for x in out if x.get("slide") == slide]
    return out

def extract_file(path: Path, max_chars: int, page: int | None, sheet: str | None, slide: int | None) -> dict[str, Any]:
    blocks = filter_blocks(blocks_for(path), page, sheet, slide)
    chunks = []
    used = 0
    truncated = False
    for block in blocks:
        text = block.get("text", "")
        prefix_parts = []
        for key in ("page", "sheet", "slide", "row", "table", "line"):
            if key in block:
                prefix_parts.append(f"{key}={block[key]}")
        prefix = "[" + " ".join(prefix_parts) + "] " if prefix_parts else ""
        chunk = prefix + text
        if used + len(chunk) + 1 > max_chars:
            remaining = max_chars - used
            if remaining > 0:
                chunks.append(chunk[:remaining])
            truncated = True
            break
        chunks.append(chunk)
        used += len(chunk) + 1
    return {
        "ok": True,
        "command": "extract",
        "path": str(path),
        "extension": path.suffix.lower(),
        "filters": {"page": page, "sheet": sheet, "slide": slide},
        "block_count": len(blocks),
        "returned_chars": sum(len(x) for x in chunks),
        "truncated": truncated,
        "text": "\n".join(chunks),
    }

def search_file(path: Path, query: str, limit: int, context_chars: int) -> dict[str, Any]:
    q = query.casefold()
    matches = []
    for block in blocks_for(path):
        text = block.get("text", "")
        lower = text.casefold()
        start = 0
        while True:
            pos = lower.find(q, start)
            if pos < 0:
                break
            a = max(0, pos - context_chars)
            b = min(len(text), pos + len(query) + context_chars)
            meta = {k: v for k, v in block.items() if k != "text"}
            matches.append({**meta, "context": text[a:b], "match_offset": pos})
            if len(matches) >= limit:
                break
            start = pos + max(1, len(query))
        if len(matches) >= limit:
            break
    return {"ok": True, "command": "search", "path": str(path), "query": query, "count": len(matches), "results": matches}

def compare_files(path_a: Path, path_b: Path, max_diff_lines: int) -> dict[str, Any]:
    text_a = plain_text(path_a)
    text_b = plain_text(path_b)
    lines_a = text_a.splitlines()
    lines_b = text_b.splitlines()
    ratio = difflib.SequenceMatcher(None, text_a, text_b).ratio()
    diff = list(difflib.unified_diff(lines_a, lines_b, fromfile=str(path_a), tofile=str(path_b), lineterm=""))
    truncated = len(diff) > max_diff_lines
    return {
        "ok": True,
        "command": "compare",
        "path_a": str(path_a),
        "path_b": str(path_b),
        "sha256_a": sha256(path_a),
        "sha256_b": sha256(path_b),
        "identical_binary": sha256(path_a) == sha256(path_b),
        "text_similarity": round(ratio, 6),
        "text_chars_a": len(text_a),
        "text_chars_b": len(text_b),
        "diff_line_count": len(diff),
        "diff_truncated": truncated,
        "diff": diff[:max_diff_lines],
    }

def make_backup(path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    dest_dir = BACKUP_ROOT / stamp
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / path.name
    shutil.copy2(path, dest)
    return dest

def replace_in_paragraph_runs(paragraph, replacements: list[tuple[str, str]]) -> tuple[int, list[str]]:
    changed = 0
    unresolved = []
    full = paragraph.text
    for old, new in replacements:
        if old not in full:
            continue
        local_changed = 0
        for run in paragraph.runs:
            if old in run.text:
                n = run.text.count(old)
                run.text = run.text.replace(old, new)
                changed += n
                local_changed += n
        if local_changed == 0 and old in full:
            unresolved.append(old)
    return changed, unresolved

def replace_docx(path: Path, replacements: list[tuple[str, str]], output: Path) -> dict[str, Any]:
    doc = Document(str(path))
    changed = 0
    unresolved = []
    for p in doc.paragraphs:
        n, u = replace_in_paragraph_runs(p, replacements)
        changed += n
        unresolved.extend(u)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    n, u = replace_in_paragraph_runs(p, replacements)
                    changed += n
                    unresolved.extend(u)
    for sec in doc.sections:
        for collection in (sec.header.paragraphs, sec.footer.paragraphs):
            for p in collection:
                n, u = replace_in_paragraph_runs(p, replacements)
                changed += n
                unresolved.extend(u)
    if changed:
        doc.save(str(output))
    return {"replacements_made": changed, "unresolved_cross_run_terms": sorted(set(unresolved))}

def replace_pptx(path: Path, replacements: list[tuple[str, str]], output: Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    changed = 0
    unresolved = []
    for slide in prs.slides:
        for shape in slide.shapes:
            text_frame = getattr(shape, "text_frame", None)
            if text_frame is not None:
                for p in text_frame.paragraphs:
                    full = p.text
                    for old, new in replacements:
                        if old not in full:
                            continue
                        local = 0
                        for run in p.runs:
                            if old in run.text:
                                n = run.text.count(old)
                                run.text = run.text.replace(old, new)
                                changed += n
                                local += n
                        if local == 0:
                            unresolved.append(old)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            full = p.text
                            for old, new in replacements:
                                if old not in full:
                                    continue
                                local = 0
                                for run in p.runs:
                                    if old in run.text:
                                        n = run.text.count(old)
                                        run.text = run.text.replace(old, new)
                                        changed += n
                                        local += n
                                if local == 0:
                                    unresolved.append(old)
    if changed:
        prs.save(str(output))
    return {"replacements_made": changed, "unresolved_cross_run_terms": sorted(set(unresolved))}

def replace_xlsx(path: Path, replacements: list[tuple[str, str]], output: Path, include_formulas: bool) -> dict[str, Any]:
    wb = load_workbook(str(path), read_only=False, data_only=False)
    changed = 0
    changed_cells = []
    skipped_formula_cells = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if not isinstance(v, str):
                        continue
                    if v.startswith("=") and not include_formulas:
                        if any(old in v for old, _ in replacements):
                            skipped_formula_cells.append(f"{ws.title}!{cell.coordinate}")
                        continue
                    new_v = v
                    for old, new in replacements:
                        if old in new_v:
                            changed += new_v.count(old)
                            new_v = new_v.replace(old, new)
                    if new_v != v:
                        cell.value = new_v
                        changed_cells.append(f"{ws.title}!{cell.coordinate}")
        if changed:
            wb.save(str(output))
    finally:
        wb.close()
    return {"replacements_made": changed, "changed_cells": changed_cells[:500], "changed_cell_count": len(changed_cells), "skipped_formula_cells": skipped_formula_cells[:200]}

def replace_textlike(path: Path, replacements: list[tuple[str, str]], output: Path) -> dict[str, Any]:
    text, enc = read_textlike(path)
    changed = 0
    new_text = text
    for old, new in replacements:
        changed += new_text.count(old)
        new_text = new_text.replace(old, new)
    if path.suffix.lower() == ".json" and changed:
        json.loads(new_text)
    if changed:
        output.write_text(new_text, encoding="utf-8")
    return {"replacements_made": changed, "source_encoding": enc, "output_encoding": "utf-8"}

def replace_file(path: Path, replacements_raw: str, output_raw: str | None, include_formulas: bool) -> dict[str, Any]:
    ext = path.suffix.lower()
    if ext not in EDITABLE:
        raise EngineError(f"Safe editing is not supported for {ext}; read/compare/search only")

    try:
        parsed = json.loads(replacements_raw)
    except json.JSONDecodeError as exc:
        raise EngineError(f"Invalid replacements JSON: {exc}") from exc
    if not isinstance(parsed, list) or not parsed:
        raise EngineError("replacements must be a non-empty JSON list")
    replacements = []
    for item in parsed:
        if not isinstance(item, dict) or "old" not in item or "new" not in item:
            raise EngineError("Each replacement must be an object with old and new")
        old, new = str(item["old"]), str(item["new"])
        if not old:
            raise EngineError("Replacement old value cannot be empty")
        replacements.append((old, new))

    output = Path(output_raw).expanduser() if output_raw else path
    if output.suffix.lower() != ext:
        raise EngineError("Output extension must match input extension")
    output.parent.mkdir(parents=True, exist_ok=True)

    before_hash = sha256(path)
    backup = make_backup(path) if output.resolve() == path.resolve() else None
    temp = output.with_name(output.name + ".document-engine.tmp") if output.resolve() == path.resolve() else output

    try:
        if ext == ".docx":
            result = replace_docx(path, replacements, temp)
        elif ext == ".pptx":
            result = replace_pptx(path, replacements, temp)
        elif ext == ".xlsx":
            result = replace_xlsx(path, replacements, temp, include_formulas)
        elif ext in TEXTLIKE:
            result = replace_textlike(path, replacements, temp)
        else:
            raise EngineError(f"Editing not implemented for {ext}")

        changed = int(result.get("replacements_made", 0))
        if changed and temp != output:
            os.replace(temp, output)
        elif not changed and temp.exists() and temp != output:
            temp.unlink(missing_ok=True)

        after_hash = sha256(output) if output.exists() else before_hash
        return {
            "ok": True,
            "command": "replace",
            "path": str(path),
            "output": str(output),
            "backup": str(backup) if backup else None,
            "before_sha256": before_hash,
            "after_sha256": after_hash,
            "changed": changed > 0,
            **result,
        }
    except Exception:
        if temp.exists() and temp != output:
            temp.unlink(missing_ok=True)
        raise

def health() -> dict[str, Any]:
    versions = {}
    import importlib.metadata as md
    for package in ("python-docx", "pypdf", "PyMuPDF", "openpyxl", "python-pptx", "xlrd"):
        try:
            versions[package] = md.version(package)
        except Exception:
            versions[package] = None
    return {
        "ok": True,
        "command": "health",
        "supported_extensions": sorted(SUPPORTED),
        "editable_extensions": sorted(EDITABLE),
        "read_only_extensions": sorted(SUPPORTED - EDITABLE),
        "dependencies": versions,
        "backup_root": str(BACKUP_ROOT),
        "safe_edit_rules": {
            "pdf": "read_only",
            "xls": "read_only",
            "docx_pptx": "same-run replacements only; cross-run matches reported unresolved",
            "xlsx": "cell-value replacements preserve workbook styles; formulas skipped by default",
            "json": "post-edit JSON validity required",
        },
    }

def main() -> int:
    parser = argparse.ArgumentParser(prog="document-engine")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("health")

    p = sub.add_parser("inspect")
    p.add_argument("path")

    p = sub.add_parser("extract")
    p.add_argument("path")
    p.add_argument("--max-chars", type=int, default=30000)
    p.add_argument("--page", type=int)
    p.add_argument("--sheet")
    p.add_argument("--slide", type=int)

    p = sub.add_parser("search")
    p.add_argument("path")
    p.add_argument("query")
    p.add_argument("--limit", type=int, default=20)
    p.add_argument("--context-chars", type=int, default=120)

    p = sub.add_parser("compare")
    p.add_argument("path_a")
    p.add_argument("path_b")
    p.add_argument("--max-diff-lines", type=int, default=80)

    p = sub.add_parser("replace")
    p.add_argument("path")
    p.add_argument("replacements_json")
    p.add_argument("--output")
    p.add_argument("--include-formulas", action="store_true")

    args = parser.parse_args()
    try:
        if args.command == "health":
            result = health()
        elif args.command == "inspect":
            result = inspect_file(ensure_path(args.path))
        elif args.command == "extract":
            result = extract_file(ensure_path(args.path), max(1000, min(args.max_chars, 200000)), args.page, args.sheet, args.slide)
        elif args.command == "search":
            result = search_file(ensure_path(args.path), args.query, max(1, min(args.limit, 100)), max(20, min(args.context_chars, 1000)))
        elif args.command == "compare":
            result = compare_files(ensure_path(args.path_a), ensure_path(args.path_b), max(10, min(args.max_diff_lines, 500)))
        elif args.command == "replace":
            result = replace_file(ensure_path(args.path), args.replacements_json, args.output, args.include_formulas)
        else:
            raise EngineError("Unknown command")
        emit(result)
        return 0
    except Exception as exc:
        emit({"ok": False, "command": args.command, "error": type(exc).__name__, "message": str(exc)})
        return 1

if __name__ == "__main__":
    raise SystemExit(main())